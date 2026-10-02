import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.report_tools import ReportTools


def query(columns, rows, truncated=False):
    return {"columns": columns, "rows": rows, "truncated": truncated}


@pytest.fixture
def report_tools(tmp_path):
    return ReportTools(output_dir=tmp_path)


def test_empty_result_exports_headers_but_cannot_be_charted(report_tools):
    artifact = report_tools.export(query(["nombre"], []), "csv")
    assert artifact.path.read_bytes().startswith(b"\xef\xbb\xbf")
    with artifact.path.open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream)) == [["nombre"]]
    with pytest.raises(ValueError, match="vacío"):
        report_tools.chart(query(["nombre", "valor"], []), {"kind": "bar", "x": "nombre", "y": "valor", "title": "T"})


def test_csv_documents_supported_mixed_types_and_neutralizes_formulas(report_tools):
    data = query(
        ["texto", "decimal", "fecha", "nulo", "binario"],
        [["=HYPERLINK(\"bad\")", Decimal("12.340"), date(2026, 1, 2), None, b"\x00\xff"]],
    )
    artifact = report_tools.export(data, "csv", "Mezcla")
    with artifact.path.open(encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream))
    assert records[1] == ["'=HYPERLINK(\"bad\")", "12.340", "2026-01-02", "", "base64:AP8="]


def test_xlsx_has_filter_format_and_preserves_dangerous_or_precise_values(report_tools):
    artifact = report_tools.export(
        query(["texto", "entero", "momento"], [["+SUM(1,1)", 12345678901234567, datetime(2026, 1, 2, 3, 4, 5)]]),
        "xlsx",
    )
    workbook = load_workbook(artifact.path)
    sheet = workbook["Datos"]
    assert sheet["A2"].value == "'+SUM(1,1)"
    assert sheet["B2"].value == "12345678901234567"
    assert sheet.auto_filter.ref == "A1:C2"
    assert sheet.freeze_panes == "A2"
    assert artifact.warnings


@pytest.mark.parametrize(
    "data,message",
    [
        (query(["x", "x"], [[1, 2]]), "duplicados"),
        (query(["x"], [[object()]]), "Tipo"),
        ({"columns": ["x"], "rows": [], "truncated": False, "path": "/tmp/a"}, "exactamente"),
    ],
)
def test_invalid_query_data_is_rejected(report_tools, data, message):
    with pytest.raises(ValueError, match=message):
        report_tools.export(data, "csv")


def test_unknown_chart_columns_and_unsuitable_values_are_rejected(report_tools):
    with pytest.raises(ValueError, match="deben existir"):
        report_tools.chart(query(["x", "y"], [[1, 2]]), {"kind": "line", "x": "missing", "y": "y", "title": "T"})
    with pytest.raises(ValueError, match="columna y"):
        report_tools.chart(query(["x", "y"], [[1, "two"]]), {"kind": "bar", "x": "x", "y": "y", "title": "T"})


def test_paths_are_generated_unique_and_inside_output_dir(report_tools, tmp_path):
    first = report_tools.export(query(["x"], [[1]]), "csv", "../../fuera")
    second = report_tools.export(query(["x"], [[1]]), "csv", "../../fuera")
    assert first.path.parent == tmp_path.resolve()
    assert first.path != second.path
    assert first.path.name.startswith("fuera_")


def test_limits_reject_instead_of_silently_dropping_rows(tmp_path):
    config = tmp_path / "reports.yml"
    config.write_text(
        "reports:\n  output_dir: out\n  max_export_rows: 1\n  max_chart_points: 1\n  max_columns: 2\n  max_cell_chars: 100\n",
        encoding="utf-8",
    )
    tools = ReportTools(config_path=config)
    with pytest.raises(ValueError, match="reducción explícita"):
        tools.export(query(["x"], [[1], [2]]), "csv")


def test_truncation_is_explicit_in_artifacts(report_tools):
    csv_artifact = report_tools.export(query(["x"], [[1]], truncated=True), "csv")
    assert csv_artifact.truncated and "TRUNCADO" in csv_artifact.path.name
    assert "total general" in csv_artifact.warnings[0]
    xlsx_artifact = report_tools.export(query(["x"], [[1]], truncated=True), "xlsx")
    assert "TRUNCADO" in load_workbook(xlsx_artifact.path).properties.description


def test_chart_embeds_plotly_for_offline_use_and_escapes_external_text(report_tools):
    artifact = report_tools.chart(
        query(["<x>", "y"], [["a", Decimal("1.5")]]),
        {"kind": "bar", "x": "<x>", "y": "y", "title": "</script><script>alert(1)</script>"},
    )
    html = artifact.path.read_text(encoding="utf-8")
    assert artifact.kind == "html"
    assert "plotly.js" in html
    assert "https://cdn.plot.ly" not in html
    assert "</script><script>alert(1)</script>" not in html
