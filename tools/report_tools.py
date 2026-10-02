"""Deterministic exports and charts for already-fetched query results.

This module deliberately has no database or model dependency.  ``QueryData`` is
the small, lossless boundary between query execution and reporting.
"""

from __future__ import annotations

import base64
import csv
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, TypedDict
from uuid import uuid4

import yaml


DEFAULT_CONFIG = Path(__file__).resolve().parents[1] / "config" / "reports.yml"
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_SAFE_NAME = re.compile(r"[^\w.-]+", re.UNICODE)
_SCALAR_TYPES = (str, int, float, bool, Decimal, date, datetime, time, bytes, type(None))


class QueryData(TypedDict):
    columns: list[str]
    rows: list[tuple[Any, ...] | list[Any]]
    truncated: bool


@dataclass(frozen=True)
class Artifact:
    """A generated file plus the facts the UI must display about it."""

    kind: str
    path: Path
    title: str
    row_count: int
    truncated: bool
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class _Settings:
    output_dir: Path
    max_export_rows: int
    max_chart_points: int
    max_columns: int
    max_cell_chars: int


class ReportTools:
    """Create files only from validated ``QueryData`` supplied by the caller."""

    def __init__(self, output_dir: str | Path | None = None, config_path: str | Path | None = None):
        config_file = Path(config_path) if config_path is not None else DEFAULT_CONFIG
        try:
            raw = yaml.safe_load(config_file.read_text(encoding="utf-8"))
        except (OSError, yaml.YAMLError) as exc:
            raise ValueError(f"No se pudo leer la configuración de reportes: {config_file}") from exc
        if not isinstance(raw, dict) or set(raw) != {"reports"} or not isinstance(raw["reports"], dict):
            raise ValueError("El YAML debe contener únicamente la sección reports")
        cfg = raw["reports"]
        expected = {"output_dir", "max_export_rows", "max_chart_points", "max_columns", "max_cell_chars"}
        if set(cfg) != expected:
            raise ValueError(f"Claves de configuración inválidas; se esperan: {sorted(expected)}")
        selected = Path(output_dir) if output_dir is not None else Path(cfg["output_dir"])
        if not selected.is_absolute():
            selected = (config_file.resolve().parent / selected).resolve()
        else:
            selected = selected.resolve()
        for key in expected - {"output_dir"}:
            if type(cfg[key]) is not int or cfg[key] < 1:
                raise ValueError(f"reports.{key} debe ser un entero positivo")
        selected.mkdir(parents=True, exist_ok=True)
        if not selected.is_dir():
            raise ValueError("output_dir no es un directorio")
        self._settings = _Settings(selected, cfg["max_export_rows"], cfg["max_chart_points"], cfg["max_columns"], cfg["max_cell_chars"])

    @property
    def output_dir(self) -> Path:
        return self._settings.output_dir

    def export(self, data: dict, format: str, title: str = "Reporte") -> Artifact:
        columns, rows, truncated = self._validate_data(data, self._settings.max_export_rows)
        if format not in {"csv", "xlsx"}:
            raise ValueError("format debe ser 'csv' o 'xlsx'")
        title = self._validate_text(title, "title")
        path = self._new_path(title, format, truncated)
        warnings: list[str] = []
        try:
            if format == "csv":
                self._write_csv(path, columns, rows)
            else:
                warnings = self._write_xlsx(path, columns, rows, title, truncated)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        if truncated:
            warnings.insert(0, "Resultado truncado: el archivo contiene sólo las filas recibidas; no se calculó un total general.")
        return Artifact(format, path, title, len(rows), truncated, tuple(warnings))

    def chart(self, data: dict, spec: dict) -> Artifact:
        columns, rows, truncated = self._validate_data(data, self._settings.max_chart_points)
        if not rows:
            raise ValueError("No se puede graficar un resultado vacío")
        if not isinstance(spec, dict) or set(spec) != {"kind", "x", "y", "title"}:
            raise ValueError("spec debe contener exactamente kind, x, y y title")
        kind = spec["kind"]
        if kind not in {"bar", "line", "scatter"}:
            raise ValueError("kind debe ser bar, line o scatter")
        x_name = self._validate_text(spec["x"], "x")
        y_name = self._validate_text(spec["y"], "y")
        title = self._validate_text(spec["title"], "title")
        if x_name not in columns or y_name not in columns:
            raise ValueError("Las columnas x e y deben existir en QueryData")
        xi, yi = columns.index(x_name), columns.index(y_name)
        x_values, y_values = [r[xi] for r in rows], [r[yi] for r in rows]
        self._validate_chart_values(x_values, y_values, kind)

        import plotly.graph_objects as go

        trace_type = go.Scatter if kind in {"line", "scatter"} else go.Bar
        kwargs: dict[str, Any] = {"x": x_values, "y": y_values}
        if kind in {"line", "scatter"}:
            kwargs["mode"] = "lines+markers" if kind == "line" else "markers"
        figure = go.Figure(data=[trace_type(**kwargs)])
        suffix = " — RESULTADO TRUNCADO" if truncated else ""
        figure.update_layout(title={"text": title + suffix}, xaxis_title=x_name, yaxis_title=y_name)
        path = self._new_path(title, "html", truncated)
        try:
            figure.write_html(path, include_plotlyjs=True, full_html=True, auto_open=False)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        warnings = (("Resultado truncado: la gráfica muestra sólo los puntos recibidos; no representa un total general.",) if truncated else ())
        return Artifact("html", path, title, len(rows), truncated, warnings)

    def _validate_data(self, data: dict, limit: int) -> tuple[list[str], list[tuple[Any, ...] | list[Any]], bool]:
        if not isinstance(data, dict) or set(data) != {"columns", "rows", "truncated"}:
            raise ValueError("QueryData debe contener exactamente columns, rows y truncated")
        columns, rows, truncated = data["columns"], data["rows"], data["truncated"]
        if not isinstance(columns, list) or not columns or len(columns) > self._settings.max_columns:
            raise ValueError("columns debe ser una lista no vacía dentro del límite configurado")
        if any(not isinstance(c, str) or not c.strip() or len(c) > self._settings.max_cell_chars for c in columns):
            raise ValueError("Cada nombre de columna debe ser texto no vacío y de tamaño permitido")
        if len(set(columns)) != len(columns):
            raise ValueError("QueryData contiene nombres de columna duplicados")
        if not isinstance(rows, list):
            raise ValueError("rows debe ser una lista")
        if len(rows) > limit:
            raise ValueError(f"El resultado excede el límite configurado de {limit}; solicite una reducción explícita")
        if type(truncated) is not bool:
            raise ValueError("truncated debe ser booleano")
        for number, row in enumerate(rows, 1):
            if not isinstance(row, (list, tuple)) or len(row) != len(columns):
                raise ValueError(f"La fila {number} no coincide con las columnas")
            for value in row:
                if not isinstance(value, _SCALAR_TYPES):
                    raise ValueError(f"Tipo de celda no admitido: {type(value).__name__}")
                if isinstance(value, str) and len(value) > self._settings.max_cell_chars:
                    raise ValueError("Una celda excede max_cell_chars")
                if isinstance(value, bytes) and len("base64:") + 4 * ((len(value) + 2) // 3) > self._settings.max_cell_chars:
                    raise ValueError("Una celda binaria excede max_cell_chars al representarse en Base64")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("NaN e infinito no son valores admitidos")
        return columns, rows, truncated

    @staticmethod
    def _validate_text(value: Any, name: str) -> str:
        if not isinstance(value, str) or not value.strip() or "\x00" in value or len(value) > 32767:
            raise ValueError(f"{name} debe ser texto no vacío")
        return value

    def _new_path(self, title: str, extension: str, truncated: bool) -> Path:
        stem = _SAFE_NAME.sub("_", title.strip()).strip("._")[:50] or "reporte"
        marker = "_TRUNCADO" if truncated else ""
        path = self.output_dir / f"{stem}{marker}_{uuid4().hex}.{extension}"
        resolved = path.resolve()
        if resolved.parent != self.output_dir or resolved.exists():
            raise RuntimeError("No se pudo generar una ruta segura y única")
        return resolved

    @staticmethod
    def _text_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, bytes):
            return "base64:" + base64.b64encode(value).decode("ascii")
        if isinstance(value, (datetime, date, time)):
            return value.isoformat()
        if isinstance(value, bool):
            return "true" if value else "false"
        return str(value)

    @classmethod
    def _excel_safe_text(cls, value: str) -> str:
        return "'" + value if value.startswith(_FORMULA_PREFIXES) else value

    def _write_csv(self, path: Path, columns: list[str], rows: list[list[Any] | tuple[Any, ...]]) -> None:
        with path.open("x", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, dialect="excel", lineterminator="\r\n")
            writer.writerow([self._excel_safe_text(c) for c in columns])
            writer.writerows([self._excel_safe_text(self._text_value(v)) for v in row] for row in rows)

    def _write_xlsx(self, path: Path, columns: list[str], rows: list[list[Any] | tuple[Any, ...]], title: str, truncated: bool) -> list[str]:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws.append([self._excel_safe_text(c) for c in columns])
        warnings: list[str] = []
        precision_warning = False
        for row in rows:
            output = []
            for value in row:
                if isinstance(value, str):
                    value = self._excel_safe_text(value)
                elif isinstance(value, bytes):
                    value = self._text_value(value)
                elif isinstance(value, (Decimal, int)) and not isinstance(value, bool) and self._exceeds_excel_precision(value):
                    value = self._excel_safe_text(str(value))
                    precision_warning = True
                elif isinstance(value, time):
                    value = value.isoformat()
                output.append(value)
            ws.append(output)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
        for index, column in enumerate(columns, 1):
            contents = [self._text_value(r[index - 1]) for r in rows]
            ws.column_dimensions[get_column_letter(index)].width = min(60, max(10, len(column) + 2, *(len(v) + 2 for v in contents)))
        wb.properties.title = title
        wb.properties.description = "RESULTADO TRUNCADO; sin total general" if truncated else "Resultado completo recibido"
        ws.oddFooter.center.text = "RESULTADO TRUNCADO — sin total general" if truncated else ""
        if precision_warning:
            warnings.append("Números con más de 15 dígitos se representaron como texto para conservar precisión en Excel.")
        wb.save(path)
        return warnings

    @staticmethod
    def _exceeds_excel_precision(value: Decimal | int) -> bool:
        """Excel retains at most 15 significant decimal digits in numeric cells."""
        if isinstance(value, int):
            return len(str(abs(value))) > 15
        return len(value.as_tuple().digits) > 15

    @staticmethod
    def _validate_chart_values(x_values: list[Any], y_values: list[Any], kind: str) -> None:
        numeric = (int, float, Decimal)
        if any(v is None or isinstance(v, bool) or not isinstance(v, numeric) for v in y_values):
            raise ValueError("La columna y debe contener únicamente números no nulos")
        if kind == "scatter" and any(v is None or isinstance(v, bool) or not isinstance(v, numeric) for v in x_values):
            raise ValueError("La columna x de scatter debe contener únicamente números no nulos")
        allowed_x = (str, int, float, Decimal, date, datetime, time)
        if kind != "scatter" and any(v is None or isinstance(v, bool) or not isinstance(v, allowed_x) for v in x_values):
            raise ValueError("La columna x contiene valores no aptos para la gráfica")
